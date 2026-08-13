"""Design system dùng chung cho các bộ slide Gladiators.

Tách khỏi ``build_slides.py`` để hai deck (kiến trúc và proposal) không giữ hai
bản token/helper song song rồi lệch nhau — đúng thứ mà chính hệ thống này tồn
tại để chống. Module cố ý KHÔNG có side effect ở cấp module.
"""

from __future__ import annotations

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

# --- design tokens ----------------------------------------------------------
NAVY = RGBColor(0x0F, 0x2A, 0x43)
INK = RGBColor(0x14, 0x1B, 0x24)
BODY = RGBColor(0x33, 0x41, 0x4E)
MUTED = RGBColor(0x6B, 0x7A, 0x8A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PANEL = RGBColor(0xF2, 0xF5, 0xF8)
LINE = RGBColor(0xD8, 0xE0, 0xE8)

TEAL = RGBColor(0x0E, 0x7C, 0x86)
BLUE = RGBColor(0x1E, 0x4D, 0x8C)
AMBER = RGBColor(0xC9, 0x6F, 0x1E)
RED = RGBColor(0xA6, 0x33, 0x3F)
PURPLE = RGBColor(0x6B, 0x46, 0x8F)
GREEN = RGBColor(0x2F, 0x7A, 0x4F)
GREY = RGBColor(0x8A, 0x97, 0xA5)

SERIES = [BLUE, TEAL, AMBER, PURPLE, RED, GREY]

FONT = "Segoe UI"
W, H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.62)
CONTENT_TOP = Inches(1.62)


def _txt(frame, text, size, color, bold=False, italic=False, space_after=0,
         align=PP_ALIGN.LEFT, para=None, font=FONT):
    p = frame.paragraphs[0] if para is None else para
    p.alignment = align
    p.space_after = Pt(space_after)
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.name = font
    return p


def box(slide, x, y, w, h, fill=None, line=None, line_w=1.0):
    from pptx.enum.shapes import MSO_SHAPE

    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    sh.shadow.inherit = False
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(line_w)
    return sh


def tb(slide, x, y, w, h):
    t = slide.shapes.add_textbox(x, y, w, h)
    tf = t.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = W, H
        self.blank = self.prs.slide_layouts[6]
        self.n = 0

    def _bg(self, slide, color=WHITE):
        box(slide, 0, 0, W, H, fill=color)

    def _chrome(self, slide, kicker):
        """Footer: kicker trái, số slide phải."""
        self.n += 1
        y = H - Inches(0.46)
        box(slide, MARGIN, y - Inches(0.08), W - 2 * MARGIN, Emu(9525), fill=LINE)
        f = tb(slide, MARGIN, y, Inches(9), Inches(0.3))
        _txt(f, kicker, 9, MUTED)
        f2 = tb(slide, W - MARGIN - Inches(1.2), y, Inches(1.2), Inches(0.3))
        _txt(f2, f"{self.n:02d}", 9, MUTED, bold=True, align=PP_ALIGN.RIGHT)

    def slide(self, title, sub=None, kicker="Gladiators · Analytical QA System"):
        s = self.prs.slides.add_slide(self.blank)
        self._bg(s)
        box(s, MARGIN, Inches(0.58), Inches(0.055), Inches(0.42), fill=TEAL)
        f = tb(s, MARGIN + Inches(0.22), Inches(0.52), W - 2 * MARGIN - Inches(0.3), Inches(0.55))
        _txt(f, title, 25, INK, bold=True)
        if sub:
            f2 = tb(s, MARGIN + Inches(0.22), Inches(1.06), W - 2 * MARGIN - Inches(0.3), Inches(0.4))
            _txt(f2, sub, 12, MUTED)
        self._chrome(s, kicker)
        return s


# ---------------------------------------------------------------------------
# helpers cho các khối nội dung hay dùng
# ---------------------------------------------------------------------------

def bullets(slide, x, y, w, items, size=13, gap=9, bullet_color=TEAL):
    """items: list[str] hoặc list[tuple[str, str]] = (đầu đề đậm, phần còn lại)."""
    tf = tb(slide, x, y, w, Inches(4.6))
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        head, rest = (it, None) if isinstance(it, str) else it
        r = p.add_run()
        r.text = "▪  "
        r.font.size = Pt(size)
        r.font.color.rgb = bullet_color
        r.font.name = FONT
        r2 = p.add_run()
        r2.text = head
        r2.font.size = Pt(size)
        r2.font.bold = rest is not None
        r2.font.color.rgb = INK if rest is not None else BODY
        r2.font.name = FONT
        if rest:
            r3 = p.add_run()
            r3.text = " " + rest
            r3.font.size = Pt(size)
            r3.font.color.rgb = BODY
            r3.font.name = FONT
    return tf


def card(slide, x, y, w, h, title, lines, accent=TEAL, title_size=12, body_size=10.5):
    box(slide, x, y, w, h, fill=PANEL)
    box(slide, x, y, Inches(0.045), h, fill=accent)
    tf = tb(slide, x + Inches(0.22), y + Inches(0.16), w - Inches(0.4), h - Inches(0.3))
    _txt(tf, title, title_size, INK, bold=True, space_after=5)
    for ln in lines:
        p = tf.add_paragraph()
        _txt(tf, ln, body_size, BODY, space_after=3, para=p)
    return tf


def stat(slide, x, y, w, value, label, accent=TEAL, vsize=27):
    tf = tb(slide, x, y, w, Inches(1.0))
    _txt(tf, value, vsize, accent, bold=True, space_after=1)
    p = tf.add_paragraph()
    _txt(tf, label, 10, MUTED, space_after=0, para=p)


def table(slide, x, y, w, rows, col_w=None, head_fill=NAVY, size=10.5,
          row_h=Inches(0.32), head_size=10.5, aligns=None):
    nrow, ncol = len(rows), len(rows[0])
    shp = slide.shapes.add_table(nrow, ncol, x, y, w, row_h * nrow)
    t = shp.table
    t.first_row = True
    if col_w:
        total = sum(col_w)
        for i, cw in enumerate(col_w):
            t.columns[i].width = Emu(int(w * cw / total))
    for ri, row in enumerate(rows):
        t.rows[ri].height = row_h
        for ci, val in enumerate(row):
            c = t.cell(ri, ci)
            c.margin_left = c.margin_right = Inches(0.09)
            c.margin_top = c.margin_bottom = Inches(0.02)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.fill.solid()
            if ri == 0:
                c.fill.fore_color.rgb = head_fill
            else:
                c.fill.fore_color.rgb = WHITE if ri % 2 else PANEL
            tf = c.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = (aligns[ci] if aligns else PP_ALIGN.LEFT)
            r = p.add_run()
            r.text = str(val)
            r.font.size = Pt(head_size if ri == 0 else size)
            r.font.bold = ri == 0
            r.font.color.rgb = WHITE if ri == 0 else BODY
            r.font.name = FONT
    return t


def mono(slide, x, y, w, h, lines, fill=RGBColor(0x10, 0x1C, 0x28), fg=RGBColor(0xC8, 0xD6, 0xE2),
         size=10, highlight=None):
    """Khối code/trace. highlight: dict[int index dòng] -> RGBColor."""
    box(slide, x, y, w, h, fill=fill)
    tf = tb(slide, x + Inches(0.2), y + Inches(0.14), w - Inches(0.35), h - Inches(0.26))
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(2)
        r = p.add_run()
        r.text = ln
        r.font.size = Pt(size)
        r.font.name = "Consolas"
        r.font.color.rgb = (highlight or {}).get(i, fg)
        r.font.bold = i in (highlight or {})
    return tf


def style_chart(chart, size=10, legend=True, legend_pos=XL_LEGEND_POSITION.BOTTOM):
    chart.font.size = Pt(size)
    chart.font.name = FONT
    chart.font.color.rgb = BODY
    chart.has_legend = legend
    if legend:
        chart.legend.position = legend_pos
        chart.legend.include_in_layout = False


def color_points(plot, colors):
    for i, pt in enumerate(plot.series[0].points):
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = colors[i % len(colors)]


# ===========================================================================
#  BUILD
# ===========================================================================
